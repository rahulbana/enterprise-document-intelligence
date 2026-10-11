# Name: Rahul Singh Bana
# Email: bana.saa@gmail.com
# LinkedIn: https://www.linkedin.com/in/rahulsinghbana
#
# One test per registered parser, built from real files constructed with the
# same library that parses them (not hand-crafted bytes) wherever a writer
# library exists. Exceptions, both documented at their own parser:
#   - MsgParser: no .msg writer library exists at all.
#   - PstParser: no .pst writer library exists either, and a real .pst needs
#     real Outlook — only the readpst-missing error path is unit-testable
#     here (the happy path was verified live against the real stack; see
#     PLAN.md Phase 10).

import io
import shutil
import tempfile
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path

import docx
import openpyxl
import pillow_heif
import pytest
import xlwt
from ebooklib import epub
from odf.draw import Frame, TextBox
from odf.draw import Page as DrawPage
from odf.opendocument import OpenDocumentPresentation, OpenDocumentSpreadsheet, OpenDocumentText
from odf.presentation import Notes
from odf.style import MasterPage, PageLayout
from odf.table import Table, TableCell, TableRow
from odf.text import H
from odf.text import P as OdfP
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from app.ingestion.parsers.base import ParsedElement
from app.ingestion.parsers.csv_parser import CsvParser, TsvParser
from app.ingestion.parsers.docx_parser import DocxParser
from app.ingestion.parsers.eml_parser import EmlParser
from app.ingestion.parsers.epub_parser import EpubParser
from app.ingestion.parsers.factory import get_parser
from app.ingestion.parsers.html_parser import HtmlParser
from app.ingestion.parsers.image_parser import ImageParser
from app.ingestion.parsers.json_parser import JsonParser
from app.ingestion.parsers.jsonl_parser import JsonlParser
from app.ingestion.parsers.markdown_parser import MarkdownParser
from app.ingestion.parsers.mbox_parser import MboxParser
from app.ingestion.parsers.mhtml_parser import MhtmlParser
from app.ingestion.parsers.msg_parser import MsgParser
from app.ingestion.parsers.odp_parser import OdpParser
from app.ingestion.parsers.ods_parser import OdsParser
from app.ingestion.parsers.odt_parser import OdtParser
from app.ingestion.parsers.pdf_parser import PdfParser
from app.ingestion.parsers.pptx_parser import PptxParser
from app.ingestion.parsers.pst_parser import PstParser
from app.ingestion.parsers.rtf_parser import RtfParser
from app.ingestion.parsers.txt_parser import TxtParser
from app.ingestion.parsers.xhtml_parser import XhtmlParser
from app.ingestion.parsers.xls_parser import XlsParser
from app.ingestion.parsers.xlsx_parser import XlsxParser
from app.ingestion.parsers.xml_parser import XmlParser
from app.ingestion.parsers.yaml_parser import YamlParser

_TESSERACT_AVAILABLE = shutil.which("tesseract") is not None
_POPPLER_AVAILABLE = shutil.which("pdftoppm") is not None


def _joined(elements: list[ParsedElement]) -> str:
    return "\n".join(element.content for element in elements)


def _text_image_bytes(text: str, *, size: tuple[int, int] = (600, 150)) -> bytes:
    image = Image.new("RGB", size, color="white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=40)
    draw.text((10, 40), text, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer.getvalue()


# --- Core documents ----------------------------------------------------


def test_txt_parser_decodes_utf8() -> None:
    elements = TxtParser().extract_elements(io.BytesIO(b"hello, world"))
    assert _joined(elements) == "hello, world"


def test_txt_parser_replaces_invalid_bytes_instead_of_raising() -> None:
    elements = TxtParser().extract_elements(io.BytesIO(b"valid \xff\xfe invalid"))
    text = _joined(elements)
    assert "valid" in text
    assert "invalid" in text


def test_markdown_parser_keeps_formatting() -> None:
    elements = MarkdownParser().extract_elements(io.BytesIO(b"# Heading\n\nSome **bold** text."))
    text = _joined(elements)
    assert "# Heading" in text
    assert "**bold**" in text


def test_docx_parser_separates_body_table_header_footer() -> None:
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "42"
    section = document.sections[0]
    section.header.paragraphs[0].text = "Confidential"
    section.footer.paragraphs[0].text = "Page footer"
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    elements = DocxParser().extract_elements(buffer)
    by_kind = {e.metadata["kind"]: e for e in elements}

    assert "First paragraph." in by_kind["body"].content
    assert "Second paragraph." in by_kind["body"].content
    assert "Name | Score" in by_kind["table"].content
    assert by_kind["table"].metadata["table_index"] == 0
    assert by_kind["header"].content == "Confidential"
    assert by_kind["footer"].content == "Page footer"


def test_rtf_parser_extracts_plain_text() -> None:
    rtf = rb"{\rtf1\ansi Hello \b bold\b0  world.}"
    elements = RtfParser().extract_elements(io.BytesIO(rtf))
    assert "Hello" in _joined(elements)
    assert "world." in _joined(elements)


def test_odt_parser_sections_body_under_heading() -> None:
    document = OpenDocumentText()
    document.text.addElement(H(outlinelevel=1, text="Chapter One"))
    document.text.addElement(OdfP(text="Body paragraph."))
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    elements = OdtParser().extract_elements(buffer)
    assert len(elements) == 1
    assert elements[0].metadata == {"heading": "Chapter One", "level": 1}
    assert "Chapter One" in elements[0].content
    assert "Body paragraph." in elements[0].content


def test_html_parser_sections_under_headings() -> None:
    html = b"<html><body><h1>Intro</h1><p>Hello</p><h2>Details</h2><p>World</p></body></html>"
    elements = HtmlParser().extract_elements(io.BytesIO(html))
    assert [e.metadata["heading"] for e in elements] == ["Intro", "Details"]
    assert "Hello" in elements[0].content
    assert "World" in elements[1].content


def test_xhtml_parser_uses_same_heading_logic() -> None:
    xhtml = b"<html><body><h1>Title</h1><p>Content here.</p></body></html>"
    elements = XhtmlParser().extract_elements(io.BytesIO(xhtml))
    assert elements[0].metadata["heading"] == "Title"
    assert "Content here." in elements[0].content


def test_mhtml_parser_extracts_html_part_from_mime_envelope() -> None:
    message = MIMEMultipart("related")
    message.attach(MIMEText("plain fallback", "plain"))
    message.attach(MIMEText("<html><body><h1>Report</h1><p>Q3 numbers</p></body></html>", "html"))

    elements = MhtmlParser().extract_elements(io.BytesIO(message.as_bytes()))
    assert elements[0].metadata["heading"] == "Report"
    assert "Q3 numbers" in elements[0].content


def test_xml_parser_one_element_per_top_level_child() -> None:
    xml = b"<catalog><book>Dune</book><book>Foundation</book></catalog>"
    elements = XmlParser().extract_elements(io.BytesIO(xml))
    assert [e.content for e in elements] == ["Dune", "Foundation"]
    assert elements[0].metadata == {"path": "/catalog/book", "index": 0}


def test_epub_parser_extracts_chapters_via_html_sectioning() -> None:
    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Test Book")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Chapter One", file_name="chap1.xhtml", lang="en")
    chapter.content = "<html><body><h1>Chapter One</h1><p>It was a dark night.</p></body></html>"
    book.add_item(chapter)
    book.toc = (chapter,)
    book.spine = [chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())

    with tempfile.TemporaryDirectory() as tmp_dir:
        path = Path(tmp_dir) / "book.epub"
        epub.write_epub(str(path), book)
        data = path.read_bytes()

    elements = EpubParser().extract_elements(io.BytesIO(data))
    matching = [e for e in elements if e.metadata.get("chapter_name") == "chap1.xhtml"]
    assert matching[0].metadata["heading"] == "Chapter One"
    assert "It was a dark night." in matching[0].content


def test_json_parser_one_element_per_top_level_key() -> None:
    elements = JsonParser().extract_elements(io.BytesIO(b'{"a": "first", "b": {"nested": 1}}'))
    assert [e.metadata["key"] for e in elements] == ["a", "b"]
    assert elements[0].content == "first"
    assert "nested" in elements[1].content


def test_jsonl_parser_one_element_per_line() -> None:
    data = b'{"id": 1}\n{"id": 2}\n'
    elements = JsonlParser().extract_elements(io.BytesIO(data))
    assert [e.metadata["line"] for e in elements] == [1, 2]
    assert "1" in elements[0].content
    assert "2" in elements[1].content


def test_yaml_parser_one_element_per_top_level_key() -> None:
    yaml_bytes = b"name: Widget\nprice: 9.99\n"
    elements = YamlParser().extract_elements(io.BytesIO(yaml_bytes))
    assert {e.metadata["key"] for e in elements} == {"name", "price"}


# --- Presentations -------------------------------------------------------


def test_pptx_parser_extracts_title_and_body_text() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Slide Title"
    slide.placeholders[1].text = "Body content here."
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)

    elements = PptxParser().extract_elements(buffer)
    by_kind = {e.metadata["kind"]: e for e in elements}

    assert by_kind["title"].content == "Slide Title"
    assert by_kind["title"].metadata["slide"] == 1
    assert "Body content here." in by_kind["textbox"].content


def test_pptx_parser_extracts_table_with_dimensions() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(
        rows=2, cols=2, left=0, top=0, width=2000000, height=1000000
    )
    table = graphic_frame.table
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "42"
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)

    elements = PptxParser().extract_elements(buffer)
    table_element = next(e for e in elements if e.metadata["kind"] == "table")

    assert "Name | Score" in table_element.content
    assert "Alice | 42" in table_element.content
    assert table_element.metadata["rows"] == 2
    assert table_element.metadata["columns"] == 2


def test_pptx_parser_extracts_notes_image_chart_and_hyperlink() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    image_buffer = io.BytesIO()
    Image.new("RGB", (10, 10), color="red").save(image_buffer, format="PNG")
    image_buffer.seek(0)
    slide.shapes.add_picture(image_buffer, 0, 0, width=900000, height=900000)

    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Revenue", (10, 20))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 1000000, 3000000, 2000000, chart_data)

    textbox = slide.shapes.add_textbox(0, 2000000, 900000, 900000)
    run = textbox.text_frame.paragraphs[0].add_run()
    run.text = "Click here"
    run.hyperlink.address = "https://example.com"

    slide.notes_slide.notes_text_frame.text = "Speaker notes for slide 1."

    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)

    elements = PptxParser().extract_elements(buffer)
    by_kind = {e.metadata["kind"]: e for e in elements}

    assert "[Image:" in by_kind["image"].content
    assert "left" in by_kind["image"].metadata["position"]
    assert "[Chart:" in by_kind["chart"].content
    assert "Revenue" in by_kind["chart"].content
    assert by_kind["textbox"].metadata["hyperlinks"] == ["https://example.com"]
    assert by_kind["notes"].content == "Speaker notes for slide 1."


def test_odp_parser_extracts_body_and_notes_per_slide() -> None:
    document = OpenDocumentPresentation()
    layout = PageLayout(name="PL1")
    document.automaticstyles.addElement(layout)
    master = MasterPage(name="Default", pagelayoutname=layout)
    document.masterstyles.addElement(master)

    page = DrawPage(name="Slide1", masterpagename=master)
    frame = Frame(width="10cm", height="5cm", x="1cm", y="1cm")
    textbox = TextBox()
    textbox.addElement(OdfP(text="Hello Slide"))
    frame.addElement(textbox)
    page.addElement(frame)

    notes = Notes()
    notes_frame = Frame(width="10cm", height="2cm", x="1cm", y="7cm")
    notes_textbox = TextBox()
    notes_textbox.addElement(OdfP(text="Speaker notes here"))
    notes_frame.addElement(notes_textbox)
    notes.addElement(notes_frame)
    page.addElement(notes)

    document.presentation.addElement(page)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    elements = OdpParser().extract_elements(buffer)
    by_kind = {e.metadata["kind"]: e for e in elements}

    assert by_kind["body"].metadata["slide"] == 1
    assert "Hello Slide" in by_kind["body"].content
    assert by_kind["notes"].content == "Speaker notes here"


# --- Spreadsheets ---------------------------------------------------------


def test_xlsx_parser_preserves_headers_rows_and_coordinates() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Inventory"
    sheet.append(["Item", "Quantity"])
    sheet.append(["Widgets", 10])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    elements = XlsxParser().extract_elements(buffer)
    assert len(elements) == 1
    element = elements[0]
    assert element.metadata["sheet"] == "Inventory"
    assert element.metadata["headers"] == ["Item", "Quantity"]
    assert "# Inventory" in element.content
    assert "Row 2 (A2:B2): Item=Widgets, Quantity=10" in element.content


def test_xlsx_parser_captures_formula_verbatim() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Calc"
    sheet.append(["A", "B", "Total"])
    sheet.append([1, 2, "=A2+B2"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    element = XlsxParser().extract_elements(buffer)[0]

    # openpyxl only caches Excel-computed results; a file saved by openpyxl
    # itself (never opened in real Excel) has none, so display falls back
    # to the formula string — the formula itself is always in metadata.
    assert element.metadata["formulas"]["C2"] == "=A2+B2"
    assert "Total=" in element.content


def test_xlsx_parser_skips_empty_sheets() -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "Empty"
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    assert XlsxParser().extract_elements(buffer) == []


def test_xls_parser_reads_legacy_binary_format() -> None:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Inventory")
    sheet.write(0, 0, "Item")
    sheet.write(0, 1, "Quantity")
    sheet.write(1, 0, "Widgets")
    sheet.write(1, 1, 10)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    element = XlsParser().extract_elements(buffer)[0]
    assert element.metadata["sheet"] == "Inventory"
    assert "Item=Widgets" in element.content


def test_ods_parser_reads_open_document_spreadsheet() -> None:
    document = OpenDocumentSpreadsheet()
    table = Table(name="Inventory")
    header_row = TableRow()
    for value in ("Item", "Quantity"):
        cell = TableCell()
        cell.addElement(OdfP(text=value))
        header_row.addElement(cell)
    table.addElement(header_row)
    data_row = TableRow()
    for value in ("Widgets", "10"):
        cell = TableCell()
        cell.addElement(OdfP(text=value))
        data_row.addElement(cell)
    table.addElement(data_row)
    document.spreadsheet.addElement(table)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    element = OdsParser().extract_elements(buffer)[0]
    assert element.metadata["sheet"] == "Inventory"
    assert "Item=Widgets, Quantity=10" in element.content


def test_csv_parser_labels_rows_with_headers() -> None:
    csv_bytes = b"Name,Score\nAlice,42\nBob,7\n"
    elements = CsvParser().extract_elements(io.BytesIO(csv_bytes))
    assert elements[0].metadata["headers"] == ["Name", "Score"]
    assert "Row 2: Name=Alice, Score=42" in elements[0].content
    assert "Row 3: Name=Bob, Score=7" in elements[0].content


def test_tsv_parser_uses_tab_delimiter() -> None:
    tsv_bytes = b"Name\tScore\nAlice\t42\n"
    elements = TsvParser().extract_elements(io.BytesIO(tsv_bytes))
    assert "Row 2: Name=Alice, Score=42" in elements[0].content


# --- Images ---------------------------------------------------------------


@pytest.mark.skipif(not _TESSERACT_AVAILABLE, reason="requires the tesseract system binary")
def test_image_parser_ocrs_text_from_a_png() -> None:
    image_bytes = _text_image_bytes("OCRTEST123")
    elements = ImageParser().extract_elements(io.BytesIO(image_bytes))
    assert "OCRTEST123" in _joined(elements).upper().replace(" ", "")


@pytest.mark.skipif(not _TESSERACT_AVAILABLE, reason="requires the tesseract system binary")
def test_image_parser_reads_heic_via_pillow_heif() -> None:
    heic_buffer = io.BytesIO()
    heif_file = pillow_heif.from_pillow(Image.new("RGB", (600, 150), color="white"))
    heif_file.save(heic_buffer, format="HEIF")
    heic_buffer.seek(0)
    # No text in the image, just confirms the HEIC codec path decodes without raising.
    elements = ImageParser().extract_elements(heic_buffer)
    assert elements == []


# --- PDF --------------------------------------------------------------------


def test_pdf_parser_extracts_real_text_directly() -> None:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(300, 150))
    pdf.drawString(20, 100, "Hello from a real PDF.")
    pdf.save()
    buffer.seek(0)

    elements = PdfParser().extract_elements(buffer)
    assert "Hello from a real PDF." in _joined(elements)
    assert elements[0].metadata == {"page": 1, "kind": "text"}


@pytest.mark.skipif(
    not (_TESSERACT_AVAILABLE and _POPPLER_AVAILABLE),
    reason="requires the tesseract and poppler (pdftoppm) system binaries",
)
def test_pdf_parser_ocrs_a_scanned_page() -> None:
    image_bytes = _text_image_bytes("SCANNEDTEXT")
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(600, 150))
    pdf.drawImage(ImageReader(io.BytesIO(image_bytes)), 0, 0, width=600, height=150)
    pdf.save()
    buffer.seek(0)

    elements = PdfParser().extract_elements(buffer)
    assert "SCANNEDTEXT" in _joined(elements).upper().replace(" ", "")


# --- Email ------------------------------------------------------------------


def test_eml_parser_extracts_headers_and_body() -> None:
    message = MIMEText("Meeting moved to 3pm.")
    message["Subject"] = "Schedule change"
    message["From"] = "alice@example.com"
    message["To"] = "bob@example.com"

    elements = EmlParser().extract_elements(io.BytesIO(message.as_bytes()))
    by_kind = {e.metadata["kind"]: e for e in elements}

    assert "Schedule change" in by_kind["header"].content
    assert "alice@example.com" in by_kind["header"].content
    assert by_kind["body"].content == "Meeting moved to 3pm."


def test_mbox_parser_walks_multiple_messages() -> None:
    import mailbox

    with tempfile.TemporaryDirectory() as tmp_dir:
        mbox_path = Path(tmp_dir) / "messages.mbox"
        box = mailbox.mbox(str(mbox_path))
        try:
            for i in range(2):
                message = MIMEText(f"Body of message {i}.")
                message["Subject"] = f"Subject {i}"
                box.add(message)
            box.flush()
        finally:
            box.close()
        data = mbox_path.read_bytes()

    elements = MboxParser().extract_elements(io.BytesIO(data))
    bodies = [e for e in elements if e.metadata["kind"] == "body"]
    assert len(bodies) == 2
    assert {e.metadata["message_index"] for e in bodies} == {0, 1}


def test_pst_parser_raises_clear_error_when_readpst_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="readpst"):
        PstParser().extract_elements(io.BytesIO(b"not a real pst"))


# --- Factory ---------------------------------------------------------------


def test_get_parser_returns_correct_type_per_extension() -> None:
    assert isinstance(get_parser("txt"), TxtParser)
    assert isinstance(get_parser("md"), MarkdownParser)
    assert isinstance(get_parser("markdown"), MarkdownParser)
    assert isinstance(get_parser("docx"), DocxParser)
    assert isinstance(get_parser("rtf"), RtfParser)
    assert isinstance(get_parser("odt"), OdtParser)
    assert isinstance(get_parser("htm"), HtmlParser)
    assert isinstance(get_parser("html"), HtmlParser)
    assert isinstance(get_parser("xhtml"), XhtmlParser)
    assert isinstance(get_parser("mhtml"), MhtmlParser)
    assert isinstance(get_parser("mht"), MhtmlParser)
    assert isinstance(get_parser("xml"), XmlParser)
    assert isinstance(get_parser("epub"), EpubParser)
    assert isinstance(get_parser("json"), JsonParser)
    assert isinstance(get_parser("jsonl"), JsonlParser)
    assert isinstance(get_parser("yaml"), YamlParser)
    assert isinstance(get_parser("yml"), YamlParser)
    assert isinstance(get_parser("pdf"), PdfParser)
    assert isinstance(get_parser("pptx"), PptxParser)
    assert isinstance(get_parser("ppsx"), PptxParser)
    assert isinstance(get_parser("odp"), OdpParser)
    assert isinstance(get_parser("xlsx"), XlsxParser)
    assert isinstance(get_parser("xlsm"), XlsxParser)
    assert isinstance(get_parser("xls"), XlsParser)
    assert isinstance(get_parser("csv"), CsvParser)
    assert isinstance(get_parser("tsv"), TsvParser)
    assert isinstance(get_parser("ods"), OdsParser)
    assert isinstance(get_parser("png"), ImageParser)
    assert isinstance(get_parser("jpg"), ImageParser)
    assert isinstance(get_parser("jpeg"), ImageParser)
    assert isinstance(get_parser("tif"), ImageParser)
    assert isinstance(get_parser("tiff"), ImageParser)
    assert isinstance(get_parser("bmp"), ImageParser)
    assert isinstance(get_parser("webp"), ImageParser)
    assert isinstance(get_parser("gif"), ImageParser)
    assert isinstance(get_parser("heic"), ImageParser)
    assert isinstance(get_parser("eml"), EmlParser)
    assert isinstance(get_parser("msg"), MsgParser)
    assert isinstance(get_parser("mbox"), MboxParser)
    assert isinstance(get_parser("pst"), PstParser)


def test_get_parser_raises_for_unknown_extension() -> None:
    with pytest.raises(ValueError, match="No parser registered"):
        get_parser("xyz")


def test_get_parser_raises_for_skipped_legacy_binary_formats() -> None:
    with pytest.raises(ValueError, match="No parser registered"):
        get_parser("doc")
    with pytest.raises(ValueError, match="No parser registered"):
        get_parser("ppt")
