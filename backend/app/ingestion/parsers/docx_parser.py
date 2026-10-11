# Name: Rahul Singh Bana
# Email: bana.saa@gmail.com
# LinkedIn: https://www.linkedin.com/in/rahulsinghbana
#
# Body paragraphs are one element (docx has no reliable page-boundary
# concept via python-docx, unlike PDF/PPTX/XLSX); tables and non-empty
# headers/footers are separate elements so they don't get flattened into
# running body text.

from typing import BinaryIO

import docx

from app.ingestion.parsers.base import ParsedElement, Parser


class DocxParser(Parser):
    def extract_elements(self, file: BinaryIO) -> list[ParsedElement]:
        document = docx.Document(file)
        elements: list[ParsedElement] = []

        body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        if body_text.strip():
            elements.append(ParsedElement(content=body_text, metadata={"kind": "body"}))

        for index, table in enumerate(document.tables):
            rows = [" | ".join(cell.text for cell in row.cells) for row in table.rows]
            content = "\n".join(rows)
            if content.strip():
                elements.append(
                    ParsedElement(content=content, metadata={"kind": "table", "table_index": index})
                )

        for section_index, section in enumerate(document.sections):
            header_text = "\n".join(p.text for p in section.header.paragraphs)
            if header_text.strip():
                elements.append(
                    ParsedElement(
                        content=header_text, metadata={"kind": "header", "section": section_index}
                    )
                )
            footer_text = "\n".join(p.text for p in section.footer.paragraphs)
            if footer_text.strip():
                elements.append(
                    ParsedElement(
                        content=footer_text, metadata={"kind": "footer", "section": section_index}
                    )
                )

        return elements
